import os
import sys
import requests
import pandas as pd
import openpyxl
from datetime import datetime, timedelta



def get_dates(start_date, end_date):
    start = datetime.strptime(start_date, "%Y-%m-%d")
    end = datetime.strptime(end_date, "%Y-%m-%d")
    return [(start + timedelta(days=x)).strftime("%Y-%m-%d") for x in range(0, (end-start).days + 1)]
   

def calculate_percentage_change(df):
    df['drop/gain'] = ((df['close'] - df['open']) / df['open']) * 100
    df['drop/gain'] = df['drop/gain'].round(2)
    df['drop/gain'] = df['drop/gain'].astype(str) + '%'
    return df

def total_percentage_change(df):
    initial_open = df.iloc[0]['open']
    final_close = df.iloc[-1]['close']
    
    total_change = ((final_close - initial_open) / initial_open * 100).round(2)
    total_change = f"{total_change}%"
    
    df['totalchange'] = [None] * (len(df) - 1) + [total_change]

    return df

def generate_graph(excel_filepath):
    from openpyxl import load_workbook
    from openpyxl.chart import LineChart, Reference

    wb = load_workbook(excel_filepath)
    ws = wb.active

    opening_price_col = 3  
    closing_price_col = 4  
    date_col = 2           

   
    chart = LineChart()
    chart.title = "Daily Opening vs Closing Price"
    #chart.y_axis.title = "Price"
    #chart.x_axis.title = "Date"

    
    opening_data = Reference(ws, min_col=opening_price_col, min_row=1, max_row=ws.max_row)
    closing_data = Reference(ws, min_col=closing_price_col, min_row=1, max_row=ws.max_row)

    
    categories = Reference(ws, min_col=date_col, min_row=2, max_row=ws.max_row)

    
    chart.add_data(opening_data, titles_from_data=True)
    chart.add_data(closing_data, titles_from_data=True)
    chart.set_categories(categories)

    chart.style = 7

    ws.add_chart(chart, "J2")
    wb.save(excel_filepath)


print("Running script...")

ticker = sys.argv[1]
excel_filename = sys.argv[2]
startdate = sys.argv[3]
enddate = sys.argv[4]

print(f"Ticker: {ticker}, Start Date: {startdate}, End Date: {enddate}, Excel Filename: {excel_filename}")

if not excel_filename.endswith('.xlsx'):
    excel_filename += '.xlsx'

api_key = "FTxQsPJCQxnUHICKYZxGuaLF1OIpWALc"
excel_filepath = r'C:\Users\owent\OneDrive\Documents\ExcelFiles\{}'.format(excel_filename)

dates = get_dates(startdate, enddate)

combined_df = pd.DataFrame()


for date in dates:
    url = f"https://api.polygon.io/v1/open-close/{ticker}/{date}?adjusted=true&apiKey={api_key}"
    print(f"Fetching data for date: {date}")

    try:
        response = requests.get(url)
        response.raise_for_status()
        data = response.json()
        df = pd.DataFrame([data])
        df = df[['symbol', 'from', 'open', 'close', 'high', 'low', 'volume']]
        
        combined_df = pd.concat([combined_df, df], ignore_index=True)

    except requests.exceptions.RequestException as e:
        print(f"Error fetching data for date {date}: {e}")
        print("Exiting script...")
        sys.exit(1)


print(excel_filepath)
combined_df = calculate_percentage_change(combined_df)


if os.path.exists(excel_filepath):
   
    print("File exists, appending data...")
    existing_df = pd.read_excel(excel_filepath)
    
    
    combined_df = pd.concat([existing_df, combined_df], ignore_index=True)
    combined_df = total_percentage_change(combined_df)
    print(combined_df.head())
    combined_df.to_excel(excel_filepath, index=False)
    generate_graph(excel_filepath)
else:
   
    print("File does not exist, writing data...")
    combined_df = total_percentage_change(combined_df)
    combined_df.to_excel(excel_filepath, index=False)
    generate_graph(excel_filepath)


print(f"Data has been written to {excel_filepath}")
